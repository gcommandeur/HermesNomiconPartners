import base64
import csv
import io
from typing import Any

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from cc_crm.core.errors import BusinessRuleError
from cc_crm.models.candidate import Candidate
from cc_crm.models.client_company import ClientCompany
from cc_crm.schemas.candidate import CandidateCreate, CandidateUpdate
from cc_crm.schemas.client_company import ClientCompanyCreate
from cc_crm.schemas.imports import (
    ImportEntity,
    ImportFilePayload,
    ImportPreview,
    ImportResult,
    ImportRow,
)
from cc_crm.services.candidates import CandidateService
from cc_crm.services.client_grouping import ClientGroupingService
from cc_crm.services.clients import ClientCompanyService

MAX_IMPORT_BYTES = 10 * 1024 * 1024

TEMPLATES: dict[str, list[str]] = {
    "candidates": [
        "candidate_number",
        "first_name",
        "last_name",
        "email",
        "phone",
        "job_title",
        "qualification",
        "postal_code",
        "city",
        "salary_expectation",
        "notice_period",
        "status",
    ],
    "clients": [
        "name",
        "status",
        "email",
        "phone",
        "website",
        "postal_code",
        "city",
        "region",
        "company_size",
        "focus",
        "fee_percentage",
        "terms_status",
    ],
    "contacts": [
        "client_company_id",
        "first_name",
        "last_name",
        "role_title",
        "email",
        "phone",
        "is_primary",
    ],
    "jobs": [
        "client_company_id",
        "title",
        "status",
        "location",
        "search_radius_km",
        "work_model",
        "salary_min",
        "salary_max",
    ],
    "submissions": [
        "candidate_id",
        "client_company_id",
        "job_mandate_id",
        "contact_person_id",
        "status",
        "summary",
        "owner",
    ],
    "tasks": [
        "candidate_id",
        "client_company_id",
        "candidate_submission_id",
        "title",
        "description",
        "due_at",
        "priority",
    ],
    "notes": ["candidate_id", "client_company_id", "candidate_submission_id", "body"],
    "sourcing_metrics": [
        "period_start",
        "period_end",
        "source",
        "campaign",
        "pitches",
        "responses",
        "positive_responses",
        "qualification_calls",
    ],
}


class ImportService:
    def __init__(self, session: Session) -> None:
        self.session = session

    def template(self, entity: ImportEntity, file_format: str) -> tuple[bytes, str, str]:
        columns = self._columns(entity)
        if file_format == "xlsx":
            workbook = Workbook()
            sheet = workbook.active
            assert sheet is not None
            sheet.title = entity
            sheet.append(columns)
            output = io.BytesIO()
            workbook.save(output)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            return output.getvalue(), f"{entity}_template.xlsx", media_type
        text_output = io.StringIO()
        csv.writer(text_output).writerow(columns)
        return text_output.getvalue().encode("utf-8-sig"), f"{entity}_template.csv", "text/csv"

    def export_candidates(self, file_format: str) -> tuple[bytes, str, str]:
        """Voll befuellter Export ALLER Kandidaten (nicht nur einer Seite) als
        Ausgangspunkt fuer eine "Master-Excel"-Bulk-Korrektur: der Nutzer
        laedt das herunter, korrigiert/ergaenzt Felder direkt in der Datei
        und laedt sie ueber denselben Import-Wizard wieder hoch - siehe
        commit() fuer die Feld-Update-Logik bei einem candidate_number-Match."""
        columns = self._columns("candidates")
        candidates = self.session.scalars(
            select(Candidate).order_by(Candidate.candidate_number)
        ).all()
        rows = [self._candidate_export_row(candidate, columns) for candidate in candidates]
        if file_format == "xlsx":
            workbook = Workbook()
            sheet = workbook.active
            assert sheet is not None
            sheet.title = "candidates"
            sheet.append(columns)
            for row in rows:
                sheet.append(row)
            output = io.BytesIO()
            workbook.save(output)
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            return output.getvalue(), "candidates_export.xlsx", media_type
        text_output = io.StringIO()
        writer = csv.writer(text_output)
        writer.writerow(columns)
        writer.writerows(rows)
        return text_output.getvalue().encode("utf-8-sig"), "candidates_export.csv", "text/csv"

    @staticmethod
    def _candidate_export_row(candidate: Candidate, columns: list[str]) -> list[Any]:
        def value_for(column: str) -> Any:
            value = getattr(candidate, column, None)
            return getattr(value, "value", value)

        return [value_for(column) for column in columns]

    def preview(self, entity: ImportEntity, payload: ImportFilePayload) -> ImportPreview:
        raw = self._decode(payload.content_base64)
        source_rows, source_columns = self._parse(payload.filename, raw)
        rows: list[ImportRow] = []
        for index, source in enumerate(source_rows, start=2):
            mapped = (
                {
                    target: source.get(source_name)
                    for source_name, target in payload.mapping.items()
                    if target
                }
                if payload.mapping
                else source
            )
            rows.append(self._validate_row(entity, index, mapped))
        return ImportPreview(
            entity=entity,
            columns=source_columns,
            rows=rows,
            valid_count=sum(not row.errors and not row.duplicate for row in rows),
            error_count=sum(bool(row.errors) for row in rows),
            duplicate_count=sum(row.duplicate for row in rows),
            commit_supported=entity in {"candidates", "clients"},
        )

    def commit(self, entity: ImportEntity, rows: list[ImportRow], confirmed: bool) -> ImportResult:
        if not confirmed:
            raise BusinessRuleError("Import must be explicitly confirmed")
        if entity not in {"candidates", "clients"}:
            raise BusinessRuleError(f"Commit for {entity} is not implemented yet")
        imported = duplicates = skipped = merged = updated = 0
        errors: list[str] = []
        is_candidates = entity == "candidates"
        service: Any = (
            CandidateService(self.session) if is_candidates else ClientCompanyService(self.session)
        )
        schema: Any = CandidateCreate if is_candidates else ClientCompanyCreate
        grouping = ClientGroupingService(self.session)
        for row in rows:
            if row.errors:
                skipped += 1
                errors.extend(f"Zeile {row.row_number}: {message}" for message in row.errors)
                continue
            # Master-Excel-Bulk-Korrektur: eine Zeile mit candidate_number,
            # die zu einem bestehenden Kandidaten passt, wird nicht als
            # Duplikat uebersprungen, sondern aktualisiert abweichende,
            # nicht-leere Felder (leere Zellen ueberschreiben nie einen
            # bestehenden Wert - siehe _candidate_field_updates). Bewusst
            # candidate_number-Match statt email (candidate_number ist der
            # stabile, eindeutige Schluessel - siehe models/candidate.py).
            if is_candidates and row.data.get("candidate_number"):
                existing_candidate = self._find_candidate_by_number(row.data["candidate_number"])
                if existing_candidate is not None:
                    changes = self._candidate_field_updates(existing_candidate, row.data)
                    if changes:
                        service.update(
                            existing_candidate.id, CandidateUpdate.model_validate(changes)
                        )
                        updated += 1
                    else:
                        duplicates += 1
                        skipped += 1
                    continue
            if row.duplicate or self._is_duplicate(entity, row.data):
                # Fuer Kanzleien wird eine erneut importierte Kanzlei (z.B.
                # aus einer weiteren Vertragsdatei derselben Kanzleigruppe)
                # nicht einfach uebersprungen - fehlende Adressfelder werden
                # in den bestehenden Datensatz nachgetragen, statt Daten zu
                # verlieren oder eine Dublette anzulegen.
                if not is_candidates and row.data.get("name"):
                    existing = grouping.find_existing_for_merge(row.data["name"])
                    if existing is not None:
                        updates = grouping.merge_address_fields(existing, row.data)
                        if updates:
                            service.repository.update(existing, updates)
                            self.session.commit()
                            merged += 1
                            continue
                duplicates += 1
                skipped += 1
                continue
            try:
                service.create(schema.model_validate(row.data))
                imported += 1
            except Exception as exception:
                skipped += 1
                errors.append(f"Zeile {row.row_number}: {exception}")
        return ImportResult(
            imported=imported,
            skipped=skipped,
            errors=errors,
            duplicates=duplicates,
            merged=merged,
            updated=updated,
        )

    def _find_candidate_by_number(self, candidate_number: str) -> Candidate | None:
        return self.session.scalar(
            select(Candidate).where(Candidate.candidate_number == candidate_number)
        )

    def _candidate_field_updates(
        self, existing: Candidate, data: dict[str, Any]
    ) -> dict[str, Any]:
        """Nur Spalten aus TEMPLATES['candidates'] (ausser candidate_number
        selbst), die in der hochgeladenen Zeile einen nicht-leeren Wert
        haben UND vom aktuellen Wert abweichen. Eine leere Zelle laesst den
        bestehenden Wert unveraendert - so kann eine unvollstaendig
        ausgefuellte Zeile nie versehentlich vorhandene gute Daten loeschen."""
        updates: dict[str, Any] = {}
        for column in TEMPLATES["candidates"]:
            if column == "candidate_number" or column not in data:
                continue
            new_value = data.get(column)
            if new_value is None or new_value == "":
                continue
            current_value = getattr(existing, column, None)
            current_comparable = getattr(current_value, "value", current_value)
            if str(new_value) != str(current_comparable if current_comparable is not None else ""):
                updates[column] = new_value
        return updates

    def _validate_row(
        self, entity: ImportEntity, row_number: int, data: dict[str, Any]
    ) -> ImportRow:
        if entity not in {"candidates", "clients"}:
            return ImportRow(row_number=row_number, data=data)
        schema = CandidateCreate if entity == "candidates" else ClientCompanyCreate
        try:
            validated = schema.model_validate(data).model_dump(mode="json")
            duplicate = self._is_duplicate(entity, validated)
            return ImportRow(row_number=row_number, data=validated, duplicate=duplicate)
        except ValidationError as exception:
            messages = [
                ".".join(str(part) for part in error["loc"]) + ": " + error["msg"]
                for error in exception.errors()
            ]
            return ImportRow(row_number=row_number, data=data, errors=messages)

    def _is_duplicate(self, entity: ImportEntity, data: dict[str, Any]) -> bool:
        if entity == "candidates":
            conditions = []
            if data.get("candidate_number"):
                conditions.append(Candidate.candidate_number == data["candidate_number"])
            if data.get("email"):
                conditions.append(Candidate.email == data["email"])
            if not conditions:
                return False
            query = select(Candidate.id).where(or_(*conditions))
            return bool(self.session.scalar(query))
        if not data.get("name"):
            return False
        query = select(ClientCompany.id).where(ClientCompany.name == data["name"])
        return bool(self.session.scalar(query))

    @staticmethod
    def _decode(value: str) -> bytes:
        try:
            raw = base64.b64decode(value, validate=True)
        except ValueError as exception:
            raise BusinessRuleError("Invalid base64 upload") from exception
        if len(raw) > MAX_IMPORT_BYTES:
            raise BusinessRuleError("Import file exceeds 10 MB")
        return raw

    @staticmethod
    def _parse(filename: str, raw: bytes) -> tuple[list[dict[str, Any]], list[str]]:
        suffix = filename.lower().rsplit(".", 1)[-1]
        if suffix == "csv":
            reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
            return list(reader), list(reader.fieldnames or [])
        if suffix == "xlsx":
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
            if sheet is None:
                raise BusinessRuleError("Uploaded workbook has no sheets")
            values = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(values)]
            return [dict(zip(headers, row, strict=False)) for row in values], headers
        raise BusinessRuleError("Only CSV and XLSX files are supported")

    @staticmethod
    def _columns(entity: ImportEntity) -> list[str]:
        if entity not in TEMPLATES:
            raise BusinessRuleError("Unknown import entity")
        return TEMPLATES[entity]
